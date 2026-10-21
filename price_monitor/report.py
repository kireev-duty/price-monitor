from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from price_monitor.storage import Change, Storage

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
RED = PatternFill("solid", fgColor="F8CBAD")
GREEN = PatternFill("solid", fgColor="C6EFCE")


def _sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    return ws


def _link(ws, row: int, col: int, url: str) -> None:
    cell = ws.cell(row=row, column=col)
    cell.hyperlink, cell.value, cell.style = url, "открыть", "Hyperlink"


def _stock(v) -> str:
    return "" if v is None else ("да" if v else "нет")


def build_report(storage: Storage, path: str | Path, changes: list[Change] | None = None) -> Path:
    """Excel-отчёт: текущие цены с изменением, изменения последнего запуска, полная история."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = _sheet(
        wb,
        "Цены",
        ["Источник", "Товар", "Ссылка", "Цена", "Было", "Изм., %", "Мин.", "Макс.", "В наличии", "Проверено"],
        [16, 50, 10, 11, 11, 10, 11, 11, 11, 20],
        first=True,
    )
    rows = storage.current()
    for r, p in enumerate(rows, start=2):
        pct = None
        if p["prev_price"] and p["price"] is not None:
            pct = round((p["price"] - p["prev_price"]) / p["prev_price"] * 100, 2)
        ws.append(
            [p["source"], p["title"], None, p["price"], p["prev_price"], pct,
             p["min_price"], p["max_price"], _stock(p["in_stock"]), p["checked_at"].replace("T", " ")]
        )
        _link(ws, r, 3, p["url"])
    if rows:
        rng = f"F2:F{len(rows) + 1}"
        ws.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=RED))
        ws.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=GREEN))
    ws.auto_filter.ref = ws.dimensions

    ws = _sheet(
        wb,
        "Изменения",
        ["Источник", "Товар", "Ссылка", "Было", "Стало", "Изм., %", "Наличие", "Событие"],
        [16, 50, 10, 11, 11, 10, 14, 14],
    )
    for r, c in enumerate(changes or [], start=2):
        pct = round(c.percent, 2) if c.percent is not None else None
        if c.is_new:
            event = "новый товар"
        elif c.old_price != c.new_price:
            event = "цена"
        else:
            event = "наличие"
        stock = _stock(c.new_in_stock) if c.is_new else f"{_stock(c.old_in_stock)} → {_stock(c.new_in_stock)}"
        ws.append([c.source, c.title, None, c.old_price, c.new_price, pct, stock, event])
        _link(ws, r, 3, c.url)
    ws.auto_filter.ref = ws.dimensions

    ws = _sheet(wb, "История", ["Источник", "Товар", "Ссылка", "Цена", "В наличии", "Дата"], [16, 50, 10, 11, 11, 20])
    for r, (source, title, url, price, in_stock, ts) in enumerate(storage.history(), start=2):
        ws.append([source, title, None, price, _stock(in_stock), ts.replace("T", " ")])
        _link(ws, r, 3, url)
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    return path
