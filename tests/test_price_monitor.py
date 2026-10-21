from datetime import datetime

import httpx
import openpyxl
import pytest

from price_monitor.config import Selectors, Settings, Source, load_config
from price_monitor.notify import format_changes
from price_monitor.report import build_report
from price_monitor.scraper import Item, crawl, parse_listing, parse_price
from price_monitor.storage import Storage

SOURCE = Source(
    name="test",
    url="https://shop.test/catalog/",
    selectors=Selectors(
        item=".card", title=".name", price=".price", link="a.name",
        availability=".stock", next_page="a.next",
    ),
)


def page(cards: str, next_href: str | None = None) -> str:
    nxt = f'<a class="next" href="{next_href}">next</a>' if next_href else ""
    return f"<html><body>{cards}{nxt}</body></html>"


def card(name: str, href: str, price: str, stock: str = "В наличии") -> str:
    return (
        f'<div class="card"><a class="name" href="{href}">{name}</a>'
        f'<span class="price">{price}</span><span class="stock">{stock}</span></div>'
    )


@pytest.mark.parametrize(
    ("text", "expected"),
    [
        ("£47.82", 47.82),
        ("1 299 ₽", 1299.0),
        ("1 299,50 руб.", 1299.5),
        ("от 2 490 руб.", 2490.0),
        ("1.299,00 €", 1299.0),
        ("$1,299.99", 1299.99),
        ("1,299", 1299.0),
        ("12,50", 12.5),
        ("Цена по запросу", None),
        ("", None),
    ],
)
def test_parse_price(text, expected):
    assert parse_price(text) == expected


def test_parse_listing_items_and_next_page():
    html = page(card("Чайник", "/p/1", "2 490 ₽") + card("Тостер", "p/2", "990 ₽", "Нет в наличии"), "?page=2")
    items, next_url = parse_listing(html, "https://shop.test/catalog/", SOURCE)
    assert [i.title for i in items] == ["Чайник", "Тостер"]
    assert items[0].url == "https://shop.test/p/1"
    assert items[1].url == "https://shop.test/catalog/p/2"
    assert [i.price for i in items] == [2490.0, 990.0]
    # «Нет в наличии» содержит «в наличии», но товар должен считаться отсутствующим
    assert [i.in_stock for i in items] == [True, False]
    assert next_url == "https://shop.test/catalog/?page=2"


def test_parse_listing_skips_broken_cards():
    html = page('<div class="card"><span class="price">100</span></div>' + card("Ок", "/ok", "100"))
    items, next_url = parse_listing(html, "https://shop.test/", SOURCE)
    assert [i.title for i in items] == ["Ок"]
    assert next_url is None


def test_crawl_follows_pages_and_respects_max_pages():
    pages = {
        "/catalog/": page(card("A", "/a", "10"), "/catalog/2"),
        "/catalog/2": page(card("B", "/b", "20"), "/catalog/3"),
        "/catalog/3": page(card("C", "/c", "30")),
    }
    transport = httpx.MockTransport(lambda req: httpx.Response(200, text=pages[req.url.path]))
    with httpx.Client(transport=transport) as client:
        settings = Settings(delay=0)
        assert [i.title for i in crawl(SOURCE, settings, client)] == ["A", "B", "C"]
        limited = Source(SOURCE.name, SOURCE.url, SOURCE.selectors, max_pages=2)
        assert [i.title for i in crawl(limited, settings, client)] == ["A", "B"]


def test_storage_detects_changes():
    st = Storage(":memory:")
    first = [Item("s", "u1", "Чайник", 100.0, True), Item("s", "u2", "Тостер", 50.0, True)]
    changes = st.save(first, datetime(2026, 10, 1, 10))
    assert all(c.is_new for c in changes) and len(changes) == 2

    assert st.save(first, datetime(2026, 10, 2, 10)) == []

    second = [Item("s", "u1", "Чайник", 90.0, True), Item("s", "u2", "Тостер", 50.0, False)]
    changes = {c.url: c for c in st.save(second, datetime(2026, 10, 3, 10))}
    assert changes["u1"].old_price == 100.0 and changes["u1"].new_price == 90.0
    assert changes["u1"].percent == pytest.approx(-10.0)
    assert changes["u2"].old_in_stock is True and changes["u2"].new_in_stock is False

    cur = {p["url"]: p for p in st.current()}
    assert cur["u1"]["price"] == 90.0 and cur["u1"]["prev_price"] == 100.0
    assert cur["u1"]["min_price"] == 90.0 and cur["u1"]["max_price"] == 100.0
    assert len(st.history()) == 6


def test_report_has_three_sheets(tmp_path):
    st = Storage(":memory:")
    st.save([Item("s", "https://x/1", "Чайник", 100.0, True)], datetime(2026, 10, 1))
    changes = st.save([Item("s", "https://x/1", "Чайник", 120.0, True)], datetime(2026, 10, 2))
    path = build_report(st, tmp_path / "r.xlsx", changes)

    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Цены", "Изменения", "История"]
    row = [c.value for c in wb["Цены"][2]]
    assert row[1] == "Чайник" and row[3] == 120.0 and row[4] == 100.0 and row[5] == 20.0
    assert wb["Цены"]["C2"].hyperlink.target == "https://x/1"
    assert wb["Изменения"]["H2"].value == "цена"
    assert wb["История"].max_row == 3


def test_format_changes_threshold_and_escaping():
    st = Storage(":memory:")
    st.save([Item("s", "u1", "A <b>", 100.0, True), Item("s", "u2", "B", 100.0, True)], datetime(2026, 10, 1))
    changes = st.save([Item("s", "u1", "A <b>", 101.0, True), Item("s", "u2", "B", 150.0, True)], datetime(2026, 10, 2))

    text = format_changes(changes)
    assert "A &lt;b&gt;" in text and "+50.0%" in text
    only_big = format_changes(changes, threshold=5)
    assert "A &lt;b&gt;" not in only_big and "B" in only_big
    assert format_changes([c for c in changes if c.url == "u1"], threshold=5) is None


def test_example_config_loads():
    config = load_config("sources.toml")
    assert len(config.sources) == 2
    assert config.sources[0].selectors.title_attr == "title"
